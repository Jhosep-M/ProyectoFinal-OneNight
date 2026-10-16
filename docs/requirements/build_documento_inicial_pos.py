from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"D:\Proyecto")
OUT = ROOT / "Proyecto-Heladeria-Monitoreo" / "docs" / "requirements" / "Documento_Inicial_POS_Heladeria_APA7.docx"
ASSETS = ROOT / ".codex" / "pos_document_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

def font(size, bold=False):
    candidates = [r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"]
    for c in candidates:
        if Path(c).exists(): return ImageFont.truetype(c, size)
    return ImageFont.load_default()

def center_text(draw, box, text, f, fill="black"):
    x1,y1,x2,y2 = box
    bb = draw.multiline_textbbox((0,0), text, font=f, align="center", spacing=5)
    draw.multiline_text(((x1+x2-(bb[2]-bb[0]))/2, (y1+y2-(bb[3]-bb[1]))/2), text, font=f, fill=fill, align="center", spacing=5)

def ellipse(draw, box, label):
    draw.ellipse(box, outline="#1F4E79", width=3, fill="#EAF3F8")
    center_text(draw, box, label, font(20))

def actor(draw, x, y, label):
    f = font(18)
    draw.ellipse((x+25,y,x+55,y+30), outline="black", width=2)
    draw.line((x+40,y+30,x+40,y+80), fill="black", width=2)
    draw.line((x+10,y+50,x+70,y+50), fill="black", width=2)
    draw.line((x+40,y+80,x+15,y+115), fill="black", width=2)
    draw.line((x+40,y+80,x+65,y+115), fill="black", width=2)
    center_text(draw, (x-25,y+122,x+105,y+160), label, f)

def make_usecase(path, title, actors, cases, connections):
    im = Image.new("RGB", (1500, 900), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((220,65,1280,830), outline="#333333", width=3)
    center_text(d, (350,74,1150,112), title, font(23, True))
    actor_positions = {}
    for label, x, y in actors:
        actor(d,x,y,label); actor_positions[label]=(x+40,y+50)
    case_positions={}
    for label, x, y, w, h in cases:
        ellipse(d,(x,y,x+w,y+h),label); case_positions[label]=(x+w/2,y+h/2)
    for a,c in connections:
        d.line((actor_positions[a],case_positions[c]),fill="#5D6D7E",width=2)
    im.save(path)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag=qn(f'w:{edge}'); el=b.find(tag)
        if el is None: el=OxmlElement(f'w:{edge}'); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:color'),'D9D9D9')

def cell_text(cell, text, bold=False, size=10):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
    r=p.add_run(text); r.bold=bold; r.font.name="Arial"; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r.font.size=Pt(size)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; borders(cell)

def setup(doc):
    s=doc.sections[0]; s.top_margin=Inches(1); s.bottom_margin=Inches(1); s.left_margin=Inches(1); s.right_margin=Inches(1)
    header = s.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    hp._p.append(fld)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); normal.font.size=Pt(12); normal.paragraph_format.line_spacing=2; normal.paragraph_format.space_after=Pt(0)
    for st in ['Title','Heading 1','Heading 2']:
        styles[st].font.name='Arial'; styles[st]._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); styles[st].font.color.rgb=RGBColor(0,0,0)
    styles['Title'].font.size=Pt(12); styles['Title'].font.bold=True
    styles['Heading 1'].font.size=Pt(12); styles['Heading 1'].font.bold=True
    styles['Heading 2'].font.size=Pt(12); styles['Heading 2'].font.bold=True

def p(doc, text='', style=None, align=None, italic=False, bold=False, first=True):
    para=doc.add_paragraph(style=style)
    para.paragraph_format.line_spacing=2
    para.paragraph_format.space_after=Pt(0)
    if first and style is None: para.paragraph_format.first_line_indent=Inches(.5)
    if align is not None: para.alignment=align
    run=para.add_run(text); run.italic=italic; run.bold=bold
    return para

def heading(doc, text, level=1): return p(doc,text,style=f'Heading {level}',first=False)

def caption(doc, text):
    q=p(doc,text,align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,first=False); q.paragraph_format.line_spacing=1

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        cell_text(t.rows[0].cells[i],h,True,9); shade(t.rows[0].cells[i],'1F4E79')
        for r in t.rows[0].cells[i].paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cell_text(cells[i],str(val),size=9)
            if idx%2==1: shade(cells[i],'F3F6F8')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    return t

def bullet(doc, text, level=0):
    q=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); q.paragraph_format.line_spacing=2; q.add_run(text)

def main():
    general=ASSETS/'casos_uso_general.png'; admin=ASSETS/'casos_uso_administracion.png'
    make_usecase(general,'Sistema POS Heladeria Cafeteria',
      [('Cajero',75,250),('Cliente',75,590),('Administrador',1320,230),('Personal de cocina',1320,590)],
      [('Abrir y cerrar\nturno',390,180,220,82),('Registrar venta\ny pagos',720,165,230,82),('Gestionar pedido\ny mesa',390,420,220,82),('Consultar puntos\ny promociones',720,415,230,82),('Preparar y actualizar\npedido',1030,560,210,82),('Anular venta o\nprocesar devolucion',700,650,270,82)],
      [('Cajero','Abrir y cerrar\nturno'),('Cajero','Registrar venta\ny pagos'),('Cajero','Gestionar pedido\ny mesa'),('Cajero','Consultar puntos\ny promociones'),('Cajero','Anular venta o\nprocesar devolucion'),('Cliente','Consultar puntos\ny promociones'),('Administrador','Anular venta o\nprocesar devolucion'),('Personal de cocina','Preparar y actualizar\npedido')])
    make_usecase(admin,'Administracion del Sistema POS',
      [('Administrador',75,350)],
      [('Gestionar productos\ny categorias',450,150,260,82),('Gestionar inventario,\ninsumos y recetas',870,150,280,82),('Gestionar clientes,\npuntos y promociones',450,410,290,82),('Gestionar usuarios\ny permisos',870,410,230,82),('Consultar reportes\ny auditoria',650,650,260,82)],
      [('Administrador','Gestionar productos\ny categorias'),('Administrador','Gestionar inventario,\ninsumos y recetas'),('Administrador','Gestionar clientes,\npuntos y promociones'),('Administrador','Gestionar usuarios\ny permisos'),('Administrador','Consultar reportes\ny auditoria')])

    doc=Document(); setup(doc)
    # APA cover
    for _ in range(7): doc.add_paragraph()
    p(doc,'Documento inicial de requisitos',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,first=False)
    p(doc,'Sistema POS para Heladeria y Cafeteria',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,first=False)
    p(doc,'Nombre de los integrantes del grupo',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
    p(doc,'Nombre de la institucion',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
    p(doc,'Asignatura',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
    p(doc,'Docente',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
    p(doc,'12 de octubre de 2026',align=WD_ALIGN_PARAGRAPH.CENTER,first=False)
    doc.add_page_break()

    heading(doc,'Indice',1)
    for item in ['1. Resumen y requisitos','1.1. Introduccion y vision general','1.2. Requisitos del sistema','1.2.1. Requisitos funcionales','1.2.2. Diagramas de casos de uso','1.2.3. Historias de usuario','1.2.4. Requisitos no funcionales','1.3. Stakeholders y usuarios','Referencias']:
        p(doc,item,first=False)
    doc.add_page_break()

    heading(doc,'1. Resumen y requisitos',1)
    p(doc,'El presente documento define los requisitos iniciales del Sistema POS para Heladeria y Cafeteria. El sistema se orienta a digitalizar y controlar las operaciones diarias de venta, caja, pedidos, mesas, inventario, productos, recetas, clientes y promociones. Su proposito es que el personal pueda atender con rapidez, mantener la trazabilidad de las operaciones y contar con informacion confiable para la administracion del negocio.')
    heading(doc,'1.1. Introduccion y vision general',2)
    p(doc,'El proyecto se denomina Sistema POS para Heladeria y Cafeteria. Es una aplicacion de punto de venta destinada a negocios que comercializan helados, bebidas, postres y productos de cafeteria. Responde a las necesidades de cajeros, personal de cocina y administradores que requieren registrar operaciones de forma ordenada y reducir errores manuales.')
    p(doc,'El objetivo general es proporcionar una plataforma segura y utilizable para gestionar el ciclo operativo del negocio, desde la apertura de caja y la toma de pedidos hasta el cierre de turno, el control de existencias y la consulta de reportes. Los objetivos especificos son agilizar el cobro, asegurar la disponibilidad de productos, registrar movimientos de inventario, administrar clientes y promociones, y conservar una auditoria de las acciones relevantes.')
    heading(doc,'Alcance del proyecto',2)
    p(doc,'El alcance incluye los modulos de ventas y pagos, caja y turnos, pedidos y mesas, productos y categorias, inventario e insumos, recetas, proveedores, clientes y puntos, promociones, devoluciones y anulaciones, usuarios y permisos, auditoria y reportes operativos.')
    heading(doc,'Fuera de alcance',2)
    p(doc,'Esta version no incluye el sistema de monitoreo de agua y energia, la recepcion o el envio de consumos, alertas ambientales, facturacion electronica ante entidades fiscales, pasarelas de pago bancarias, compras en linea, aplicacion movil para clientes ni integracion con plataformas de reparto. Estos elementos podran evaluarse como proyectos posteriores, sin formar parte de los requisitos funcionales de este documento.')

    heading(doc,'1.2. Requisitos del sistema',2)
    heading(doc,'1.2.1. Requisitos funcionales',2)
    functional=[
      ('RF-01','El sistema debe permitir iniciar sesion a usuarios autorizados y aplicar permisos segun su rol.'),
      ('RF-02','El cajero debe poder abrir un turno de caja con un monto inicial y solo mantener un turno abierto a la vez.'),
      ('RF-03','El cajero debe poder registrar ventas de uno o varios productos, validando existencia, cantidades y stock disponible.'),
      ('RF-04','El sistema debe calcular subtotales, descuentos, promociones y total de cada venta.'),
      ('RF-05','El sistema debe permitir pagos con uno o varios metodos y validar que su suma coincida con el total.'),
      ('RF-06','El sistema debe permitir crear pedidos asociados o no a una mesa y actualizar su estado: pendiente, en preparacion, listo y cobrado.'),
      ('RF-07','El personal autorizado debe poder crear, actualizar, inactivar y consultar productos y categorias.'),
      ('RF-08','El sistema debe registrar ingresos, salidas y ajustes de inventario, y advertir cuando un insumo alcance su stock minimo.'),
      ('RF-09','El sistema debe asociar recetas a productos y descontar automaticamente los insumos requeridos al concretar una venta.'),
      ('RF-10','El sistema debe permitir administrar proveedores, clientes, puntos de fidelidad y promociones vigentes.'),
      ('RF-11','El personal autorizado debe poder anular una venta o procesar devoluciones parciales o totales, con motivo, restitucion de inventario y trazabilidad.'),
      ('RF-12','El cajero debe poder cerrar su turno; el sistema debe calcular el efectivo esperado, registrar el monto real e indicar diferencias.'),
      ('RF-13','El administrador debe poder consultar reportes operativos de ventas, caja, inventario, productos y clientes.'),
      ('RF-14','El sistema debe conservar un registro de auditoria de acciones criticas, incluyendo ventas, anulaciones, devoluciones, ajustes y cierres de caja.')]
    table(doc,['ID','Requisito funcional'],functional,[.8,5.9])

    heading(doc,'1.2.2. Diagramas de casos de uso',2)
    p(doc,'Los siguientes diagramas presentan las interacciones principales entre los actores y el Sistema POS. El cliente participa de manera indirecta en la consulta de beneficios; las operaciones se registran por medio del personal autorizado.')
    doc.add_picture(str(general),width=Inches(6.45)); caption(doc,'Figura 1\nCasos de uso operativos del Sistema POS')
    doc.add_picture(str(admin),width=Inches(6.45)); caption(doc,'Figura 2\nCasos de uso de administracion del Sistema POS')

    heading(doc,'1.2.3. Historias de usuario',2)
    stories=[
      ('HU-01','Cajero','abrir mi turno con el monto inicial','iniciar las operaciones de caja de forma controlada.','Solo se permite un turno abierto por cajero.'),
      ('HU-02','Cajero','registrar una venta con varios productos y metodos de pago','cobrar correctamente al cliente.','Se valida stock y que los pagos sumen el total.'),
      ('HU-03','Cajero','crear un pedido y asociarlo a una mesa','organizar la atencion y el despacho.','El pedido conserva estados trazables hasta su cobro.'),
      ('HU-04','Personal de cocina','ver los pedidos pendientes y actualizar su estado','preparar y entregar los productos en orden.','Puede marcar pedidos en preparacion y listos.'),
      ('HU-05','Administrador','gestionar productos, categorias, insumos y recetas','mantener actualizado el catalogo y el control de costos.','No se eliminan productos con historial; se inactivan.'),
      ('HU-06','Administrador','registrar movimientos de inventario y stock minimo','evitar quiebres de existencias.','Todo movimiento indica tipo, cantidad, motivo y responsable.'),
      ('HU-07','Cajero','consultar clientes y aplicar promociones vigentes','ofrecer beneficios y fidelizar clientes.','La promocion debe estar vigente y cumplir sus reglas.'),
      ('HU-08','Cajero autorizado','anular una venta o registrar una devolucion','corregir una operacion sin perder el historial.','Se exige motivo, permiso y restitucion de inventario.'),
      ('HU-09','Cajero','cerrar mi turno indicando el efectivo real','conocer el resultado de mi caja.','El sistema calcula efectivo esperado y diferencia.'),
      ('HU-10','Administrador','consultar reportes y auditoria','tomar decisiones basadas en informacion confiable.','La informacion se filtra segun permisos.')]
    table(doc,['ID','Como','Quiero','Para','Criterio de aceptacion'],stories,[.55,.85,1.6,1.55,2.1])

    heading(doc,'1.2.4. Requisitos no funcionales',2)
    nfs=[
      ('RNF-01 Rendimiento','Las operaciones habituales de consulta, venta y pedido deben responder en un tiempo objetivo de hasta 3 segundos en condiciones normales de operacion.'),
      ('RNF-02 Seguridad','El sistema debe autenticar usuarios mediante tokens seguros, aplicar autorizacion por roles y permisos, validar entradas, proteger contra inyeccion SQL y no exponer secretos ni trazas tecnicas al usuario final.'),
      ('RNF-03 Disponibilidad','El sistema debe estar disponible durante el horario operativo. Debe contar con manejo centralizado de errores y respaldos periodicos de la base de datos.'),
      ('RNF-04 Integridad','Las ventas, pagos, anulaciones, devoluciones, inventario y cierres deben ejecutarse de forma transaccional para evitar registros parciales o inconsistentes.'),
      ('RNF-05 Escalabilidad','La arquitectura debe permitir separar frontend, backend y base de datos para atender el crecimiento de usuarios, productos y transacciones sin rediseñar los modulos principales.'),
      ('RNF-06 Usabilidad','La interfaz debe ser clara, en espanol, adaptable a equipos de escritorio y facilitar tareas frecuentes de caja con la menor cantidad razonable de pasos.'),
      ('RNF-07 Trazabilidad','Las acciones criticas deben quedar registradas con usuario, fecha, operacion realizada y datos necesarios para su revision.'),
      ('RNF-08 Cumplimiento','El tratamiento de datos personales de clientes debe limitarse a la finalidad operativa, aplicar control de acceso y respetar la normativa local aplicable de privacidad y proteccion de datos.')]
    table(doc,['ID y atributo','Especificacion'],nfs,[1.45,5.25])

    heading(doc,'1.3. Stakeholders y usuarios',2)
    p(doc,'Los stakeholders son las personas que usan, administran, se benefician o influyen en el Sistema POS. La definicion de sus responsabilidades permite asignar permisos de manera coherente y asegurar que cada modulo atienda una necesidad del negocio.')
    stakeholders=[
      ('Propietario o gerente','Define politicas comerciales, revisa resultados y autoriza decisiones relevantes.','Acceso a reportes y supervision administrativa.'),
      ('Administrador del sistema','Configura catalogos, usuarios, roles, recetas, promociones e inventario.','Acceso administrativo segun permisos asignados.'),
      ('Cajero','Abre y cierra turnos, registra ventas, pagos, pedidos y clientes.','Acceso a caja, ventas, pedidos y consultas permitidas.'),
      ('Personal de cocina','Consulta pedidos y actualiza su avance de preparacion.','Acceso limitado al modulo de pedidos.'),
      ('Encargado de inventario','Registra existencias, movimientos, proveedores e insumos.','Acceso a inventario, productos y recetas.'),
      ('Cliente','Recibe atencion, compra productos y puede acumular o consultar puntos.','No requiere acceso directo al sistema interno.')]
    table(doc,['Stakeholder o usuario','Interes o responsabilidad','Perfil y acceso'],stakeholders,[1.4,2.9,2.4])

    heading(doc,'Referencias',1)
    p(doc,'Proyecto Heladeria y Monitoreo. (2026). Requisitos funcionales POS [Documento interno del proyecto].',first=False)
    p(doc,'Proyecto Heladeria y Monitoreo. (2026). POS Heladeria Cafeteria [Documentacion tecnica interna].',first=False)
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    print(OUT)

if __name__ == '__main__': main()

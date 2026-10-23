from pathlib import Path
import re
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent

def clean(s): return re.sub(r'\s+', ' ', s).strip()

def sql_tables(folder):
    result=[]
    for f in sorted(folder.glob('*.sql')):
        text=f.read_text(encoding='utf-8', errors='ignore')
        for m in re.finditer(r'CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?(?:(\w+)\.)?(\w+)\s*\((.*?)\);', text, re.I|re.S):
            name=m.group(2); body=m.group(3); cols=[]
            for line in body.splitlines():
                x=clean(line.rstrip(','))
                if not x or x.startswith('--') or re.match(r'(CONSTRAINT|PRIMARY KEY|FOREIGN KEY|UNIQUE|CHECK|EXCLUDE)',x,re.I): continue
                mm=re.match(r'"?([a-zA-Z_][\w]*)"?\s+([\w]+(?:\s*\([^)]*\))?(?:\s+\w+)?)',x)
                if mm:
                    c,t=mm.groups(); flags=[]
                    if re.search(r'PRIMARY\s+KEY',x,re.I): flags.append('PK')
                    if re.search(r'NOT\s+NULL',x,re.I): flags.append('NN')
                    if re.search(r'REFERENCES\s+',x,re.I): flags.append('FK')
                    cols.append((c,t,'; '.join(flags) or '—'))
            if cols and name not in {x[0] for x in result}: result.append((name,cols,f.name))
    return result

pos_tables=sql_tables(ROOT/'database/pos/migrations')
mon_tables=sql_tables(ROOT/'database/monitoreo')

def shade(cell, fill='1F4E79'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell(cell,text,bold=False,color=None):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(10); r.font.name='Times New Roman'
    r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    trPr=t.rows[0]._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)
    for c,h in zip(t.rows[0].cells,headers): shade(c); set_cell(c,h,True,(255,255,255))
    for ix,row in enumerate(rows):
        cells=t.add_row().cells
        for c,v in zip(cells,row):
            if ix%2: shade(c,'EAF2F8')
            set_cell(c,v)
    doc.add_paragraph()

def add_para(doc,text,style=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.line_spacing=2.0; p.paragraph_format.space_after=Pt(6)
    r=p.add_run(text); r.font.size=Pt(12); r.font.name='Times New Roman'
    r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    return p

def add_heading(doc,text,level=1):
    p=doc.add_heading(text,level=level); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(6)
    return p

def add_portada(doc):
    for _ in range(3): doc.add_paragraph()
    
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('UNIVERSIDAD PRIVADA DOMINGO SABIO'); r.bold=True; r.font.size=Pt(16)
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('FACULTAD DE INGENIERÍA'); r.bold=True; r.font.size=Pt(14)
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('DOCUMENTO INICIAL - PARTE 2'); r.bold=True; r.font.size=Pt(18)
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Sistema de Gestión para Heladería Cafetería con Módulo Integrado de Monitoreo de Consumo de Agua y Energía'); r.font.size=Pt(14); r.italic=True
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    
    datos=[
        ('Asignatura:', 'Programación Web II'),
        ('Docente:', 'Hugo Arnaldo Guzmán Centellas'),
        ('Integrantes:', 'Airton (Product Owner)'),
        ('', 'Ana (Scrum Master)'),
        ('', 'Emerson (Development Team)'),
        ('', 'Fabio (Development Team)'),
        ('', 'Luis (Development Team)'),
        ('Fecha:', '[Hasta el lunes]')
    ]
    
    for label, value in datos:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if label:
            r=p.add_run(label + ' '); r.bold=True; r.font.size=Pt(12)
            r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
        r=p.add_run(value); r.font.size=Pt(12)
        r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')

def build_document():
    d=Document()
    sec=d.sections[0]
    sec.top_margin=Inches(1); sec.bottom_margin=Inches(1)
    sec.left_margin=Inches(1); sec.right_margin=Inches(1)
    
    normal=d.styles['Normal']
    normal.font.name='Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    normal.font.size=Pt(12)
    
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        s=d.styles[name]
        s.font.name='Times New Roman'
        s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
        s.font.color.rgb=RGBColor(0,0,0)
    
    add_portada(d)
    d.add_page_break()
    
    add_heading(d,'1.4 Asignación de Roles en la Estructura SCRUM')
    add_para(d,'El equipo de trabajo se organiza bajo la metodología SCRUM con los siguientes roles y responsabilidades:')
    
    add_table(d,['Rol','Integrante','Funciones en el Proyecto'],[
        ('Product Owner','Airton','Define requisitos, prioriza el backlog del producto, valida entregas, representa a los stakeholders, asegura el valor del producto.'),
        ('Scrum Master','Ana','Facilita las ceremonias SCRUM, elimina impedimentos, asegura el cumplimiento del proceso, coach del equipo, organiza sprint planning y dailies.'),
        ('Development Team','Emerson','Desarrollo backend, APIs REST, lógica de negocio, pruebas unitarias, integración con bases de datos.'),
        ('Development Team','Fabio','Desarrollo frontend, interfaces React, componentes UI, integración con APIs, pruebas de usuario.'),
        ('Development Team','Luis','Desarrollo backend y frontend, módulo de monitoreo, integraciones, pruebas de integración, despliegue Docker.')
    ])
    
    add_para(d,'Stakeholders: Usuarios finales (cafetería/heladería), docente, universidad.')
    add_para(d,'Administradores: Equipo de desarrollo con permisos RBAC para gestión de usuarios, roles y permisos del sistema.')
    
    d.add_page_break()
    add_heading(d,'2. Documentación Técnica')
    
    add_heading(d,'2.1 Arquitectura',2)
    add_para(d,'El sistema sigue una arquitectura de capas separada para cada aplicación (POS y Monitoreo), comunicándose mediante APIs REST autenticadas.')
    
    add_heading(d,'Stack Tecnológico',3)
    add_table(d,['Capa','Tecnología','Versión/Detalle'],[
        ('Frontend','React + Vite','React 18/19, Vite 5'),
        ('Backend','Node.js + Express','Express 5'),
        ('ORM','Sequelize','Sequelize 6'),
        ('Base de datos','PostgreSQL (Supabase)','PostgreSQL 15+'),
        ('Autenticación','Supabase Auth + JWT','JWT con refresh tokens'),
        ('Autorización','RBAC + RLS','Roles, permisos, Row Level Security'),
        ('Validación','Zod','Validación de schemas'),
        ('Seguridad','Helmet + CORS + Rate Limit','Headers, cross-origin, throttling'),
        ('Contenedores','Docker + Docker Compose','Multi-stage builds'),
        ('Orquestación','Kubernetes','Deployments, services, ingress'),
        ('Cloud','AWS','Previsto: EC2, RDS, S3, Load Balancer')
    ])
    
    add_heading(d,'Capas del Backend',3)
    add_para(d,'Estructura de capas: routes → controllers → services → repositories → models')
    add_para(d,'• Routes: Define endpoints y middlewares.')
    add_para(d,'• Controllers: Recibe la petición y devuelve la respuesta (actualmente handlers inline en POS).')
    add_para(d,'• Services: Contiene reglas de negocio.')
    add_para(d,'• Repositories: Acceso a datos.')
    add_para(d,'• Models: Modelos Sequelize y definiciones de tabla.')
    
    add_heading(d,'Flujo de Datos',3)
    add_para(d,'1. Frontend → apiFetch (JWT Bearer) → Backend')
    add_para(d,'2. Backend valida JWT → deriva userId del token')
    add_para(d,'3. Rutas validan con Zod → llaman services')
    add_para(d,'4. Services delegan a funciones PostgreSQL (SECURITY DEFINER) para operaciones críticas')
    add_para(d,'5. Funciones PG validan stock, precios, auditoría en transacción')
    
    add_heading(d,'2.2 Modelo de Base de Datos',2)
    add_para(d,'El sistema utiliza dos bases de datos independientes sin claves foráneas cruzadas. POS y Monitoreo son propietarios de sus propios datos. La integración se realiza mediante APIs REST con idempotencia.')
    
    add_heading(d,'Diccionario de Datos - POS',3)
    if pos_tables:
        for name, cols, src in pos_tables[:15]:
            add_para(d,f'Tabla: {name} (fuente: {src})',style='Heading 4')
            add_table(d,['Columna','Tipo','Restricción'],cols)
    else:
        add_para(d,'[Diccionario de datos POS disponible en migraciones SQL]')
    
    add_heading(d,'Diccionario de Datos - Monitoreo',3)
    if mon_tables:
        for name, cols, src in mon_tables[:10]:
            add_para(d,f'Tabla: {name} (fuente: {src})',style='Heading 4')
            add_table(d,['Columna','Tipo','Restricción'],cols)
    else:
        add_para(d,'[Diccionario de datos Monitoreo disponible en migraciones SQL]')
    
    add_heading(d,'2.3 Diagramas UML',2)
    add_para(d,'Los diagramas de entidad-relación muestran las relaciones entre las tablas principales de cada sistema. Se encuentran disponibles en docs/uml/diagrama-entidades.md del repositorio.')
    
    add_para(d,'Principales entidades POS:')
    add_para(d,'• Usuario, Rol, Permiso, RolPermiso (autenticación y autorización)')
    add_para(d,'• TurnoCaja, Venta, DetalleVenta, Pago (operaciones de venta)')
    add_para(d,'• Producto, Categoria, Insumo, RecetaInsumo (catálogo e inventario)')
    add_para(d,'• MovimientoInventario, Proveedor (trazabilidad de stock)')
    add_para(d,'• Cliente, MovimientoPuntos, Promocion (clientes y fidelización)')
    add_para(d,'• Mesa, Pedido, DetallePedido (gestión de mesas y pedidos)')
    add_para(d,'• ConsumoReportado, ColaIntegracion (integración con Monitoreo)')
    add_para(d,'• AuditoriaAccion (trazabilidad de acciones)')
    
    add_para(d,'Principales entidades Monitoreo:')
    add_para(d,'• Organizacion, UsuarioOrganizacion (aislamiento por organización)')
    add_para(d,'• Integracion (credenciales API con hashes)')
    add_para(d,'• PuntoMedicion, TipoRecurso (puntos de medición de agua/energía)')
    add_para(d,'• RecepcionConsumoPOS, ColaProcesamiento (recepción idempotente)')
    add_para(d,'• RegistroConsumo, UmbralClasificacion (clasificación de consumos)')
    add_para(d,'• Alerta, Notificacion, EntregaAlerta (sistema de alertas)')
    add_para(d,'• MetaReduccion, Tarifa, Recomendacion (gestión de recursos)')
    
    add_heading(d,'2.4 API',2)
    add_para(d,'Todos los endpoints requieren header Authorization: Bearer <JWT> (Supabase Auth). Base URL: /api/v1')
    
    add_heading(d,'Endpoints POS',3)
    add_table(d,['Método','Ruta','Permiso','Descripción'],[
        ('GET','/sales','venta.consultar','Listar ventas'),
        ('POST','/sales','venta.crear','Crear venta'),
        ('POST','/sales/:id/anular','venta.anular','Anular venta'),
        ('GET','/shifts','turno.consultar','Listar turnos'),
        ('POST','/shifts','turno.abrir','Abrir turno'),
        ('POST','/shifts/:id/cerrar','turno.cerrar','Cerrar turno'),
        ('GET','/orders','pedido.consultar','Listar pedidos'),
        ('POST','/orders','pedido.crear','Crear pedido'),
        ('PATCH','/orders/:id','pedido.gestionar','Actualizar pedido'),
        ('POST','/orders/:id/cobrar','pedido.gestionar','Cobrar pedido'),
        ('GET','/mesas','mesa.consultar','Listar mesas'),
        ('POST','/mesas','mesa.gestionar','Crear mesa'),
        ('GET','/products','producto.consultar','Listar productos'),
        ('POST','/products','producto.gestionar','Crear producto'),
        ('GET','/inventory/insumos','inventario.consultar','Listar insumos'),
        ('POST','/inventory/movimientos','inventario.movimiento','Registrar movimiento'),
        ('GET','/customers','cliente.consultar','Listar clientes'),
        ('POST','/payments','pago.gestionar','Crear pago'),
        ('POST','/returns','devolucion.procesar','Procesar devolución'),
        ('GET','/promotions','promocion.consultar','Listar promociones'),
        ('GET','/audit','auditoria.consultar','Listar acciones de auditoría'),
        ('POST','/integrations/alerts','API-key','Recibir alerta de Monitoreo')
    ])
    
    add_heading(d,'Endpoints de Integración',3)
    add_para(d,'POS → Monitoreo: POST /api/v1/integrations/consumption')
    add_para(d,'Payload: consumoExternoId, tipoRecurso, cantidad, unidadMedida, fechaConsumo, organizacionExternaId, origen')
    add_para(d,'Requisitos: HTTPS, autenticación, validación, idempotencia, registro de recepción, reintentos, backoff, timeout.')
    
    add_para(d,'Monitoreo → POS: POST /api/v1/integrations/alerts')
    add_para(d,'Payload: alertaId, nivel, tipoRecurso, mensaje, fechaGeneracion')
    add_para(d,'Requisitos: Autenticación, validación, idempotencia, reintentos, registro de entrega.')
    
    add_heading(d,'2.5 Tecnologías Utilizadas',2)
    add_table(d,['Categoría','Tecnología','Propósito'],[
        ('Frontend','React 18/19 + Vite','Interfaces de usuario reactivas'),
        ('Backend','Node.js + Express 5','API REST y lógica de negocio'),
        ('Base de datos','PostgreSQL (Supabase)','Almacenamiento relacional'),
        ('ORM','Sequelize 6','Mapeo objeto-relacional'),
        ('Autenticación','Supabase Auth','Gestión de usuarios y sesiones'),
        ('Seguridad','JWT + RBAC + RLS','Autorización a múltiples niveles'),
        ('Validación','Zod','Validación de schemas en runtime'),
        ('Testing','Jest + Vitest','Pruebas unitarias e integración'),
        ('Contenedores','Docker + Docker Compose','Empaquetado y orquestación local'),
        ('Orquestación','Kubernetes','Despliegue en producción'),
        ('Cloud','AWS','Infraestructura en la nube')
    ])
    
    d.add_page_break()
    add_heading(d,'3. Procedimiento de Despliegue')
    
    add_para(d,'El sistema se despliega mediante Docker Compose para desarrollo local y Kubernetes para producción.')
    
    add_heading(d,'Despliegue Local con Docker',3)
    add_para(d,'1. Instalar Docker Desktop y Docker Compose.')
    add_para(d,'2. Clonar el repositorio y navegar a la raíz del proyecto.')
    add_para(d,'3. Copiar .env.example a .env y completar las credenciales de Supabase, URLs y API keys.')
    add_para(d,'4. Ejecutar: docker compose build')
    add_para(d,'5. Ejecutar: docker compose up -d')
    add_para(d,'6. Verificar servicios: docker compose ps')
    add_para(d,'7. Consultar logs: docker compose logs')
    
    add_para(d,'Servicios expuestos:')
    add_para(d,'• POS Frontend: http://localhost:5173')
    add_para(d,'• POS Backend: http://localhost:3000')
    add_para(d,'• Monitoreo Frontend: http://localhost:5174')
    add_para(d,'• Monitoreo Backend: http://localhost:4000')
    
    add_heading(d,'Despliegue en Producción',3)
    add_para(d,'1. Elegir plataforma (Render, Railway, Koyeb o AWS).')
    add_para(d,'2. Crear cuatro servicios a partir de los Dockerfiles.')
    add_para(d,'3. Cargar secretos en el gestor de variables de la plataforma.')
    add_para(d,'4. Configurar URLs internas de POS y Monitoreo, CORS y health checks.')
    add_para(d,'5. Publicar frontends y verificar endpoints.')
    add_para(d,'6. Ejecutar la prueba integrada de 17 pasos.')
    
    d.add_page_break()
    add_heading(d,'3.1 Acta de Planificación del Sprint')
    
    add_para(d,'Sprint N°: [En blanco - completar hasta el lunes]')
    add_para(d,'Fecha: [En blanco - completar hasta el lunes]')
    
    add_heading(d,'Objetivo del Sprint',3)
    add_para(d,'[Completar con el objetivo específico del sprint, por ejemplo: "Implementar módulo completo de ventas con integración de pagos y generación de reportes"]')
    
    add_heading(d,'Historias de Usuario Seleccionadas',3)
    add_table(d,['ID','Historia de Usuario','Prioridad','Estimación (puntos)'],[
        ('HU-001','[Completar]','Alta','[X]'),
        ('HU-002','[Completar]','Media','[Y]'),
        ('HU-003','[Completar]','Alta','[Z]')
    ])
    
    add_heading(d,'Estimación Total',3)
    add_para(d,'Total de puntos del sprint: [Suma de puntos]')
    
    add_heading(d,'Riesgos Identificados',3)
    add_table(d,['Riesgo','Probabilidad','Impacto','Mitigación'],[
        ('[Riesgo 1]','[Alta/Media/Baja]','[Alto/Medio/Bajo]','[Estrategia de mitigación]'),
        ('[Riesgo 2]','[Alta/Media/Baja]','[Alto/Medio/Bajo]','[Estrategia de mitigación]')
    ])
    
    add_heading(d,'Integrantes Participantes',3)
    add_para(d,'• Airton (Product Owner) - Validación de historias y criterios de aceptación')
    add_para(d,'• Ana (Scrum Master) - Facilitación de la planificación')
    add_para(d,'• Emerson (Development Team) - Estimación técnica y asignación de tareas')
    add_para(d,'• Fabio (Development Team) - Estimación técnica y asignación de tareas')
    add_para(d,'• Luis (Development Team) - Estimación técnica y asignación de tareas')
    
    d.add_page_break()
    add_heading(d,'3.2 Registro de Scrum Diario')
    
    add_para(d,'Formato de registro diario para seguimiento del avance del sprint.')
    
    add_table(d,['Fecha','Integrante','¿Qué hice ayer?','¿Qué haré hoy?','Impedimentos'],[
        ('[Fecha]','Airton','[Completar]','[Completar]','Ninguno'),
        ('[Fecha]','Ana','[Completar]','[Completar]','Ninguno'),
        ('[Fecha]','Emerson','Diseño de BD','Crear tablas','Ninguno'),
        ('[Fecha]','Fabio','Diseño de UI','Implementar componentes','Ninguno'),
        ('[Fecha]','Luis','Configuración backend','Conectar con BD','Ninguno')
    ])
    
    add_para(d,'Responsable: Equipo de Desarrollo')
    
    d.add_page_break()
    add_heading(d,'3.3 Burn Down Chart')
    
    add_para(d,'El Burn Down Chart visualiza el avance del sprint mostrando el trabajo pendiente a lo largo del tiempo.')
    
    add_heading(d,'Tabla de Seguimiento',3)
    add_table(d,['Día','Fecha','Trabajo Pendiente (puntos)','Trabajo Completado (puntos)','Observaciones'],[
        ('1','[Fecha]','100','0','Inicio del sprint'),
        ('2','[Fecha]','90','10','[Completar]'),
        ('3','[Fecha]','75','25','[Completar]'),
        ('4','[Fecha]','60','40','[Completar]'),
        ('5','[Fecha]','45','55','[Completar]'),
        ('6','[Fecha]','30','70','[Completar]'),
        ('7','[Fecha]','15','85','[Completar]'),
        ('8','[Fecha]','0','100','Sprint completado')
    ])
    
    add_para(d,'Nota: Los valores de trabajo pendiente deben actualizarse diariamente según el progreso real del equipo.')
    
    OUT_path=OUT/'Actividad 3 SI.docx'
    d.save(OUT_path)
    print(f'Documento generado: {OUT_path}')

if __name__=='__main__':
    build_document()
